"""Member roster export: one spreadsheet row per member, joining Contacts,
the latest Enrollment, today's active Authorization, and the first
EmergencyContact on file.

Four bulk queries (not per-member lookups) keep the export fast over a
network-share database; assembling rows is pure Python so it's unit-testable.
"""
from datetime import date, datetime

from db.members import (
    _read_connection, _drop_read_connection, _access_date,
    format_phone, format_ssn, format_medicaid, format_medicare,
    decode_auth_days, parse_flexible_date,
)

COLUMNS = [
    "Center Id", "Status", "Last Name", "First Name", "Chinese Name", "DOB",
    "Health Plan", "Member ID", "Medicaid", "Medicare", "SSN", "Language",
    "Case Manager", "Home Tell", "Cell", "Address", "PCP", "Hospital",
    "Notes", "Gender", "Long Lat", "HHA",
    "Enrollment Date",
    "Auth Days", "Auth Start", "Auth End",
    "Emergency_Full Name", "Emergency_Phone Number", "Emergency_Relationship",
]

_CONTACTS_QUERY = (
    "SELECT [Center ID],[Last Name],[First Name],[Chinese Name],[DOB],"
    "[Health Plan],[Member ID],[Medicaid],[Medicare],[SSN],[Language],"
    "[Case Manager],[Home Tell],[Cell],[Address],[PCP],[Hospital],[Notes],"
    "[Gender],[Long Lat],[HHA] FROM [Contacts] "
    "ORDER BY [Center ID]"
)
_ENROLLMENTS_QUERY = ("SELECT [Center ID],[start_date],[end_date] "
                      "FROM [Enrollment]")
_AUTHS_QUERY = ("SELECT [Center ID],[auth_start],[auth_end],[auth_days] "
                "FROM [Authorization]")
_EMERGENCY_QUERY = ("SELECT [ID],[Center ID],[Full Name],[Phone Number],"
                    "[Relationship] FROM [EmergencyContact]")

def _fetch_all(db_path: str, sql: str, _retry: bool = True) -> list:
    """One bulk SELECT over the cached read connection, with the same
    stale-connection retry as get_member_context."""
    import pyodbc
    conn = _read_connection(db_path)
    try:
        c = conn.cursor()
        c.execute(sql)
        return c.fetchall()
    except pyodbc.Error:
        _drop_read_connection(db_path)
        if _retry:
            return _fetch_all(db_path, sql, _retry=False)
        raise


def format_auth_days_dotted(auth_days) -> str:
    """Stored '1,3,5' -> '1.3.5' (day numbers, period-separated, sorted)."""
    days = decode_auth_days(auth_days or "")
    return ".".join(str(d) for d in sorted(days))


def pick_active_auth(auths: list[tuple], today) -> tuple | None:
    """The (start, end, days) auth in effect today: auth_start <= today <=
    auth_end, same rule as the Auths tab's 'active' status pill (a missing
    start or end doesn't disqualify). Latest start wins; None when nothing
    is active — the spreadsheet shows blank cells so lapsed coverage is
    easy to filter."""
    active = []
    for start, end, days in auths:
        start_d, end_d = _access_date(start), _access_date(end)
        if start_d is not None and start_d > today:
            continue
        if end_d is not None and end_d < today:
            continue
        active.append((start_d, end_d, days))
    if not active:
        return None
    return max(active, key=lambda a: (a[0] or date.min))


def build_export_rows(contacts, enrollments, auths, emergency, today) -> list[list]:
    """Assemble spreadsheet rows (matching COLUMNS) from raw table rows.

    contacts:    rows in _CONTACTS_QUERY order (drives row order)
    enrollments: (center_id, start_date, end_date) — latest start becomes
                 Enrollment Date; the rows also decide Status
                 (Active/Terminated, the same latest-enrollment-has-an-end
                 rule as the sidebar's terminated marks)
    auths:       (center_id, auth_start, auth_end, auth_days)
    emergency:   (id, center_id, full_name, phone, relationship) — the row
                 with the lowest id (first on file) is used
    Pure (no DB), so it is unit-testable.
    """
    from db.members import terminated_ids_from_rows
    terminated = terminated_ids_from_rows(enrollments, today)
    enroll_by_member: dict[int, date] = {}
    for cid, start, _end in enrollments:
        if cid is None or start is None:
            continue
        start = _access_date(start)
        cid = int(cid)
        if cid not in enroll_by_member or start > enroll_by_member[cid]:
            enroll_by_member[cid] = start

    auths_by_member: dict[int, list] = {}
    for cid, start, end, days in auths:
        if cid is None:
            continue
        auths_by_member.setdefault(int(cid), []).append((start, end, days))

    emergency_by_member: dict[int, tuple] = {}
    for rec_id, cid, name, phone, rel in emergency:
        if cid is None:
            continue
        cid = int(cid)
        current = emergency_by_member.get(cid)
        if current is None or rec_id < current[0]:
            emergency_by_member[cid] = (rec_id, name, phone, rel)

    rows = []
    for r in contacts:
        if r[0] is None:            # skip Contacts rows with no Center ID
            continue
        cid = int(r[0])
        active = pick_active_auth(auths_by_member.get(cid, []), today)
        em = emergency_by_member.get(cid)
        rows.append([
            cid,
            "Terminated" if cid in terminated else "Active",
            r[1] or "", r[2] or "", r[3] or "",
            # DOB as a real date cell (MM/DD/YYYY) even when the column holds
            # text like '8/23/1953'; unparseable text passes through visibly.
            parse_flexible_date(r[4]) or (r[4] or ""),
            r[5] or "", r[6] or "",
            format_medicaid(r[7]), format_medicare(r[8]), format_ssn(r[9]),
            r[10] or "", r[11] or "",
            format_phone(r[12]), format_phone(r[13]),
            r[14] or "", r[15] or "", r[16] or "", r[17] or "",
            r[18] or "", r[19] or "", r[20] or "",
            enroll_by_member.get(cid),
            format_auth_days_dotted(active[2]) if active else "",
            active[0] if active else None,
            active[1] if active else None,
            (em[1] or "") if em else "",
            format_phone(em[2]) if em else "",
            (em[3] or "") if em else "",
        ])
    return rows


def write_members_xlsx(path: str, rows: list[list]) -> None:
    """Write COLUMNS + rows to `path`: bold frozen header, MM/DD/YYYY date
    cells, and column widths sized to their content."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "MM/DD/YYYY"

    widths = {}
    for row in [COLUMNS] + rows:
        for i, v in enumerate(row, start=1):
            text = f"{v:%m/%d/%Y}" if isinstance(v, (date, datetime)) else str(v or "")
            # Notes can be paragraphs; cap so one memo doesn't blow up a column.
            widths[i] = min(max(widths.get(i, 0), len(text) + 2), 40)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)


EXPIRING_COLUMNS = ["Center Id", "Name", "Health Plan", "Expiring Date", "Notes"]


def members_expiring_in_month(members, terminated_ids, auth_rows,
                              year: int, month: int) -> list[dict]:
    """Active members whose coverage ends inside the given month, for the
    expiring-auths report.

    A member's coverage end is the latest auth_end across their auths — the
    same rule as the notification bell — so adding a renewal that pushes the
    end past the month drops them from the report. Rows are grouped by
    health plan (then soonest end, then name) so one continuous table keeps
    each plan's members together. Pure (no DB), so it is unit-testable.
    """
    latest: dict[int, date] = {}
    for cid, end in auth_rows:
        if cid is None or end is None:
            continue
        cid = int(cid)
        end = end.date() if isinstance(end, datetime) else end
        if cid not in latest or end > latest[cid]:
            latest[cid] = end
    rows = []
    for m in members:
        cid = m["center_id"]
        if cid in terminated_ids:
            continue
        end = latest.get(cid)
        if end is None or (end.year, end.month) != (year, month):
            continue
        rows.append({
            "center_id": cid,
            "name": f"{m.get('last_name', '')}, {m.get('first_name', '')}",
            "health_plan": m.get("health_plan") or "",
            "end": end,
        })
    rows.sort(key=lambda r: (r["health_plan"].lower(), r["end"],
                             r["name"].lower()))
    return rows


def _write_record_sheet(path: str, sheet_title: str, columns: list[str],
                        data_rows: list[list], widths: tuple,
                        center_cols: set[int]) -> None:
    """A printable record sheet: every cell ruled with borders (including
    empty fill-in columns), shaded header, 22pt rows for handwriting, and
    print setup — slim margins, header repeated on every page, horizontally
    centered, fixed 100% scale (any fit-to-width shrink makes Excel render
    hairline borders at uneven weights, so column widths are chosen to
    genuinely fit the printable width instead)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    thin = Side(style="thin", color="000000")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E8E8E8")
    centered = Alignment(horizontal="center", vertical="center")
    vcenter = Alignment(vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.border = grid
        cell.fill = header_fill
        cell.alignment = centered
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    for values in data_rows:
        ws.append(values)
        row_i = ws.max_row
        ws.row_dimensions[row_i].height = 22
        for col in range(1, len(columns) + 1):
            cell = ws.cell(row=row_i, column=col)
            cell.border = grid
            cell.alignment = centered if col in center_cols else vcenter
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "MM/DD/YYYY"

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4,
                                  header=0.2, footer=0.2)
    ws.print_title_rows = "1:1"
    ws.page_setup.scale = 100
    ws.print_options.horizontalCentered = True
    wb.save(path)


def write_expiring_xlsx(path: str, rows: list[dict], month_label: str) -> None:
    """The expiring-auths record sheet (empty Notes column to work the
    renewals by hand)."""
    _write_record_sheet(
        path, f"Expiring {month_label}", EXPIRING_COLUMNS,
        [[r["center_id"], r["name"], r["health_plan"], r["end"], ""]
         for r in rows],
        widths=(9, 24, 11, 12, 42), center_cols={1, 3, 4})


BIRTHDAY_COLUMNS = ["Center Id", "Alt ID", "Name", "Birthday", "Sign", "Date"]


def members_with_birthday_in_month(members, terminated_ids,
                                   month: int) -> list[dict]:
    """Active members whose birthday falls in the given month, sorted by day
    of month (then name) so the sheet reads like a calendar. DOBs are parsed
    however the database stores them (date, ISO text, or M/D/YYYY text);
    members without a parseable DOB are skipped. Pure — unit-testable."""
    from db.members import parse_flexible_date
    rows = []
    for m in members:
        cid = m["center_id"]
        if cid in terminated_ids:
            continue
        dob = parse_flexible_date(m.get("dob"))
        if dob is None or dob.month != month:
            continue
        rows.append({
            "center_id": cid,
            # The corpus value is already the display value (decrypted at
            # load when a session password is set), so it passes through.
            "alt_id": m.get("alt_id"),
            "name": f"{m.get('last_name', '')}, {m.get('first_name', '')}",
            "dob": dob,
        })
    rows.sort(key=lambda r: (r["dob"].day, r["name"].lower()))
    return rows


def write_birthday_xlsx(path: str, rows: list[dict], month_label: str) -> None:
    """The birthdays record sheet (empty Sign and Date columns to fill in
    by hand)."""
    _write_record_sheet(
        path, f"Birthdays {month_label}", BIRTHDAY_COLUMNS,
        [[r["center_id"], r["alt_id"], r["name"], r["dob"], "", ""]
         for r in rows],
        widths=(10, 12, 28, 14, 24, 13), center_cols={1, 2, 4})


ABSENCE_COLUMNS = ["Center Id", "Name", "Leave Type", "Start Date",
                   "End Date", "Notes"]


def member_absences_in_month(members, terminated_ids, absence_rows,
                             year: int, month: int) -> list[dict]:
    """Active members' absences overlapping the given month — one row per
    absence, so a member out twice appears twice. A missing end date means
    the absence is still ongoing, so it matches every month from its start
    on. Sorted by start date then name. Pure (no DB) — unit-testable.

    absence_rows: (center_id, leave_type, start, end, notes) tuples from
    get_all_absences.
    """
    import calendar as _cal
    first = date(year, month, 1)
    last = date(year, month, _cal.monthrange(year, month)[1])
    by_id = {m["center_id"]: m for m in members}
    rows = []
    for cid, leave_type, start, end, notes in absence_rows:
        if cid is None or int(cid) in terminated_ids:
            continue
        m = by_id.get(int(cid))
        if m is None:
            continue
        start, end = _access_date(start), _access_date(end)
        if start is None or start > last or (end is not None and end < first):
            continue
        rows.append({
            "center_id": int(cid),
            "name": f"{m.get('last_name', '')}, {m.get('first_name', '')}",
            "leave_type": leave_type or "",
            "start": start,
            "end": end,
            "notes": notes or "",
        })
    rows.sort(key=lambda r: (r["start"], r["name"].lower()))
    return rows


def write_absence_xlsx(path: str, rows: list[dict], month_label: str) -> None:
    """The absences record sheet; open-ended absences leave End Date blank."""
    _write_record_sheet(
        path, f"Absences {month_label}", ABSENCE_COLUMNS,
        [[r["center_id"], r["name"], r["leave_type"], r["start"], r["end"],
          r["notes"]] for r in rows],
        widths=(9, 24, 12, 12, 12, 29), center_cols={1, 4, 5})


# ── monthly meal sheet ───────────────────────────────────────────────────
#
# A blank workbook staff fill in by hand: three identically laid-out sheets
# (breakfast, meal ticket, lunch) listing every member enrolled and
# authorized at some point in the month, one column per day. Day cells are
# empty but shaded — orange when the member is authorized that weekday,
# yellow when the center is closed, blue otherwise — so the totals at the
# bottom can split meals into authorized / not authorized. The split uses
# SUMPRODUCT against a hidden 1/0 mask sheet instead of the hand-made
# template's sum-by-cell-colour macro, so the file stays a plain .xlsx.

MEAL_SHEET_TITLES = ("会员早餐统计-B", "会员餐券统计-B", "会员午餐统计-B")
MEAL_MASK_SHEET = "_mask"
MEAL_FIXED_COLUMNS = ["ID", "Alt ID", "Name", "Health plan"]
# The template's theme colours resolved to RGB: accent2 tint 0.4, accent4
# tint 0.8, accent5 tint 0.8.
MEAL_FILL_AUTHORIZED = "F4B183"
MEAL_FILL_CLOSED = "FFF2CC"
MEAL_FILL_OPEN = "DEEBF7"
DAY_AUTHORIZED, DAY_CLOSED, DAY_NOT_AUTHORIZED = (
    "authorized", "closed", "not_authorized")


def _month_span(year: int, month: int) -> tuple[date, date]:
    import calendar as _cal
    return (date(year, month, 1),
            date(year, month, _cal.monthrange(year, month)[1]))


def members_for_meal_sheet(members, groups, enrollment_rows, auth_rows,
                           year: int, month: int,
                           alt_ids_unlocked: bool) -> list[dict]:
    """Members on the month's meal sheet: enrolled on some day of the month
    AND holding at least one authorization overlapping it, sorted by Center
    ID like the hand-made template. Month overlap replaces the "active
    today" rule the other reports use, so a past month still lists members
    terminated since. Pure (no DB) — unit-testable.

    enrollment_rows: (center_id, start, end); auth_rows: (center_id, start,
    end, auth_days) — raw bulk-query tuples. A missing enrollment start
    disqualifies the row; a missing auth start does not (as in
    pick_active_auth). groups: {center_id: Group text} for Location.
    alt_ids_unlocked: the corpus alt ids are decrypted only when a session
    password is set; when it isn't, the column is left blank.
    """
    first, last = _month_span(year, month)
    enrolled: set[int] = set()
    for cid, start, end in enrollment_rows:
        if cid is None:
            continue
        start, end = _access_date(start), _access_date(end)
        if start is None or start > last or (end is not None and end < first):
            continue
        enrolled.add(int(cid))
    auths: dict[int, list[tuple]] = {}
    for cid, start, end, days in auth_rows:
        if cid is None:
            continue
        start, end = _access_date(start), _access_date(end)
        if (start is not None and start > last) or \
           (end is not None and end < first):
            continue
        auths.setdefault(int(cid), []).append((start, end, days or ""))
    rows = []
    for m in members:
        cid = m["center_id"]
        if cid not in enrolled or cid not in auths:
            continue
        rows.append({
            "center_id": cid,
            "alt_id": m.get("alt_id") if alt_ids_unlocked else None,
            "name": f"{m.get('last_name', '')}, {m.get('first_name', '')}",
            "health_plan": m.get("health_plan") or "",
            "location": groups.get(cid, "") or "",
            "auths": auths[cid],
        })
    rows.sort(key=lambda r: r["center_id"])
    return rows


def closed_days_in_month(year: int, month: int, holidays,
                         operating_days) -> set[int]:
    """Day numbers the center is closed: company holidays plus every
    weekday without an OperatingDays row (so an empty table closes the
    whole month — that is what the table means). holidays: get_holidays
    dicts; operating_days: get_operating_days {day_of_week: row}."""
    first, last = _month_span(year, month)
    holiday_dates = {h["date"] for h in holidays if h.get("date")}
    return {d for d in range(1, last.day + 1)
            if date(year, month, d) in holiday_dates
            or date(year, month, d).isoweekday() not in operating_days}


def _safe_auth_days(value) -> set[int]:
    try:
        return decode_auth_days(value or "")
    except ValueError:      # hand-edited text like "Mon" in Access
        return set()


def meal_day_states(auths, year: int, month: int,
                    closed_days: set[int]) -> list[str]:
    """One state per day of the month (index day-1): closed wins; otherwise
    authorized when any auth covers that exact date and its auth_days
    includes that weekday (union across overlapping auths); otherwise not
    authorized. auths: (start, end, auth_days) with dates or None."""
    first, last = _month_span(year, month)
    spans = [(s, e, _safe_auth_days(days)) for s, e, days in auths]
    states = []
    for d in range(1, last.day + 1):
        day = date(year, month, d)
        if d in closed_days:
            states.append(DAY_CLOSED)
            continue
        wd = day.isoweekday()
        authorized = any(
            (s is None or s <= day) and (e is None or e >= day) and wd in days
            for s, e, days in spans)
        states.append(DAY_AUTHORIZED if authorized else DAY_NOT_AUTHORIZED)
    return states


def load_meal_sheet_inputs(db_path: str) -> dict:
    """The five bulk reads the meal sheet needs, fetched once per dialog:
    enrollments, auths, groups, holidays, operating_days."""
    from db.members import get_member_groups
    from db.company_calendar import get_holidays, get_operating_days
    return {
        "enrollments": _fetch_all(db_path, _ENROLLMENTS_QUERY),
        "auths": _fetch_all(db_path, _AUTHS_QUERY),
        "groups": get_member_groups(db_path),
        "holidays": get_holidays(db_path),
        "operating_days": get_operating_days(db_path),
    }


def write_meal_sheet_xlsx(path: str, rows: list[dict], year: int, month: int,
                          closed_days: set[int]) -> None:
    """Three identical meal sheets plus the hidden mask sheet. Layout copies
    the hand-made template: Calibri 14, thin borders on every cell (blank
    day cells included), 44/26pt rows, dates as m/d, frozen at E2, landscape
    at 79% with the header row and the ID..Health plan columns repeated on
    every printed page. With no rows only the header is written."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    first, last = _month_span(year, month)
    n_days = last.day
    day_col = lambda d: 4 + d                      # E = day 1
    total_col, loc_col = 5 + n_days, 6 + n_days
    first_day, last_day = get_column_letter(5), get_column_letter(4 + n_days)
    last_row = 1 + len(rows)                        # last member row

    thin = Side(style="thin", color="000000")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    font = Font(name="Calibri", size=14)
    bold = Font(name="Calibri", size=14, bold=True)
    centered = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    fills = {
        DAY_AUTHORIZED: PatternFill("solid", fgColor=MEAL_FILL_AUTHORIZED),
        DAY_CLOSED: PatternFill("solid", fgColor=MEAL_FILL_CLOSED),
        DAY_NOT_AUTHORIZED: PatternFill("solid", fgColor=MEAL_FILL_OPEN),
    }
    states = [meal_day_states(r["auths"], year, month, closed_days)
              for r in rows]

    def day_range(col_letter):
        return f"{col_letter}2:{col_letter}{last_row}"

    def row_total(r):
        return f"=SUM({first_day}{r}:{last_day}{r})"

    def fill_sheet(ws):
        header = (MEAL_FIXED_COLUMNS
                  + [date(year, month, d) for d in range(1, n_days + 1)]
                  + ["Total", "Location"])
        ws.append(header)
        for col in range(1, loc_col + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = bold
            cell.border = grid
            cell.alignment = centered
            if 5 <= col <= 4 + n_days:
                cell.number_format = "m/d;@"
        ws.row_dimensions[1].height = 44

        for i, r in enumerate(rows):
            row_i = i + 2
            values = ([r["center_id"], r["alt_id"], r["name"],
                       r["health_plan"]]
                      + [None] * n_days
                      + [row_total(row_i), r["location"] or None])
            ws.append(values)
            ws.row_dimensions[row_i].height = 26
            for col in range(1, loc_col + 1):
                cell = ws.cell(row=row_i, column=col)
                cell.font = font
                cell.border = grid
                cell.alignment = left if col in (3, 4) else centered
                if 5 <= col <= 4 + n_days:
                    cell.fill = fills[states[i][col - 5]]

        if rows:
            summary = [
                ("Total", lambda L: f"=SUM({day_range(L)})"),
                ("Authorized",
                 lambda L: f"=SUMPRODUCT({day_range(L)},"
                           f"'{MEAL_MASK_SHEET}'!{day_range(L)})"),
                ("Not_Authorized",
                 lambda L: f"=SUMPRODUCT({day_range(L)},"
                           f"1-'{MEAL_MASK_SHEET}'!{day_range(L)})"),
            ]
            for label, formula in summary:
                row_i = ws.max_row + 1
                ws.append([label, None, None, None]
                          + [formula(get_column_letter(day_col(d)))
                             for d in range(1, n_days + 1)]
                          + [row_total(row_i), None])
                ws.row_dimensions[row_i].height = 26
                for col in range(1, loc_col + 1):
                    cell = ws.cell(row=row_i, column=col)
                    cell.font = bold
                    cell.border = grid
                    cell.alignment = centered

        for col, width in ((1, 16), (2, 16), (3, 29), (4, 18),
                           (total_col, 9), (loc_col, 10)):
            ws.column_dimensions[get_column_letter(col)].width = width
        for d in range(1, n_days + 1):
            ws.column_dimensions[get_column_letter(day_col(d))].width = 9
        ws.freeze_panes = "E2"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.scale = 79
        ws.page_margins = PageMargins(left=0.25, right=0.25,
                                      top=0.75, bottom=0.75)
        ws.print_title_rows = "1:1"
        ws.print_title_cols = "A:D"

    wb = Workbook()
    wb.active.title = MEAL_SHEET_TITLES[0]
    fill_sheet(wb.active)
    for title in MEAL_SHEET_TITLES[1:]:
        fill_sheet(wb.create_sheet(title))

    # The mask mirrors the member rows' day cells: 1 where authorized, 0
    # otherwise (closed days count as not authorized). Same row/column
    # positions on every sheet, so one mask serves all three.
    mask = wb.create_sheet(MEAL_MASK_SHEET)
    mask.append(MEAL_FIXED_COLUMNS
                + [date(year, month, d) for d in range(1, n_days + 1)])
    for i, r in enumerate(rows):
        mask.append([r["center_id"], None, r["name"], None]
                    + [1 if s == DAY_AUTHORIZED else 0 for s in states[i]])
    mask.sheet_state = "hidden"
    wb.active = 0
    wb.save(path)


def export_members_xlsx(db_path: str, out_path: str, today=None) -> int:
    """Query, assemble, and write the roster. Returns the number of rows."""
    today = today or date.today()
    rows = build_export_rows(
        _fetch_all(db_path, _CONTACTS_QUERY),
        _fetch_all(db_path, _ENROLLMENTS_QUERY),
        _fetch_all(db_path, _AUTHS_QUERY),
        _fetch_all(db_path, _EMERGENCY_QUERY),
        today,
    )
    write_members_xlsx(out_path, rows)
    return len(rows)
