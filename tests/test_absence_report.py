"""Monthly absences report: overlap filter, xlsx, dialog."""
import os
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from db.export import (
    member_absences_in_month, write_absence_xlsx, ABSENCE_COLUMNS,
)


def member(cid, last, first):
    return {"center_id": cid, "last_name": last, "first_name": first,
            "health_plan": "HF"}


MEMBERS = [
    member(1, "Chan", "Mei"),
    member(2, "Lu", "Wei"),
    member(3, "Wong", "Ka"),
    member(4, "Li", "Jun"),        # terminated below
    member(5, "Ng", "Sam"),
]

# (center_id, leave_type, start, end, notes) — as the bulk query returns them.
ABSENCES = [
    (1, "Hospital", date(2026, 1, 28), date(2026, 2, 10), "ICU"),   # spans into Feb
    (2, "Vacation", datetime(2026, 2, 3), datetime(2026, 2, 14), None),  # inside Feb
    (3, "Personal", date(2026, 1, 5), None, ""),                    # open-ended
    (2, "Hospital", date(2026, 2, 20), date(2026, 3, 2), "recheck"),  # 2nd row, cid 2
    (5, "Vacation", date(2026, 3, 5), date(2026, 3, 9), ""),        # outside Feb
    (4, "Hospital", date(2026, 2, 1), date(2026, 2, 5), ""),        # terminated
    (None, "Hospital", date(2026, 2, 1), date(2026, 2, 5), ""),     # no Center ID
    (1, "Personal", None, None, ""),                                # no start date
]


def test_overlap_filter_sorted_by_start_then_name():
    rows = member_absences_in_month(MEMBERS, {4}, ABSENCES, 2026, 2)
    assert [(r["center_id"], r["leave_type"]) for r in rows] == [
        (3, "Personal"),        # Jan 5 open-ended, still ongoing
        (1, "Hospital"),        # Jan 28 – Feb 10
        (2, "Vacation"),        # Feb 3 – Feb 14
        (2, "Hospital"),        # Feb 20 – Mar 2
    ]
    assert rows[0]["name"] == "Wong, Ka"
    assert rows[1]["notes"] == "ICU"
    # Access datetimes come back as plain dates.
    assert rows[2]["start"] == date(2026, 2, 3)
    assert rows[2]["end"] == date(2026, 2, 14)
    assert rows[0]["end"] is None                    # open-ended stays None


def test_terminated_and_invalid_rows_excluded():
    rows = member_absences_in_month(MEMBERS, {4}, ABSENCES, 2026, 2)
    assert 4 not in [r["center_id"] for r in rows]
    # December 2025: only the open-ended absence hasn't started yet, and
    # nothing else touches the month.
    assert member_absences_in_month(MEMBERS, {4}, ABSENCES, 2025, 12) == []


def test_one_row_per_absence():
    rows = member_absences_in_month(MEMBERS, {4}, ABSENCES, 2026, 2)
    assert len([r for r in rows if r["center_id"] == 2]) == 2


def test_xlsx_round_trip(tmp_path):
    from openpyxl import load_workbook
    rows = member_absences_in_month(MEMBERS, {4}, ABSENCES, 2026, 2)
    path = tmp_path / "absences.xlsx"
    write_absence_xlsx(str(path), rows, "February 2026")

    ws = load_workbook(str(path)).active
    assert [c.value for c in ws[1]] == ABSENCE_COLUMNS
    assert ws.title == "Absences February 2026"
    first = [c.value for c in ws[2]]
    assert first[0] == 3 and first[1] == "Wong, Ka"
    assert first[2] == "Personal"
    assert first[3].date() == date(2026, 1, 5)
    assert first[4] is None                          # open-ended: empty cell
    assert ws.cell(row=2, column=4).number_format == "MM/DD/YYYY"
    # Same ruled record-sheet treatment as the other reports.
    for col in range(1, len(ABSENCE_COLUMNS) + 1):
        assert ws.cell(row=2, column=col).border.bottom.style == "thin"
    assert ws.page_setup.scale == 100
    assert ws.print_title_rows in ("1:1", "$1:$1")


def test_dialog_defaults_and_counts(qapp, monkeypatch):
    from gui import absence_report as ar
    monkeypatch.setattr("db.members.get_all_absences",
                        lambda _db: ABSENCES)
    dlg = ar.AbsenceReportDialog(
        "fake.accdb", MEMBERS, {4}, today=date(2026, 2, 15))
    assert dlg._month.currentText() == "February"    # defaults to this month
    assert dlg._year.value() == 2026
    assert "4 absences" in dlg._count.text()
    assert dlg._btn_save.isEnabled()

    dlg._month.setCurrentIndex(11)                   # December 2026:
    assert "1 absence " in dlg._count.text()         # only the open-ended one

    dlg._year.setValue(2025)                         # December 2025: nothing
    assert "0 absences" in dlg._count.text()
    assert not dlg._btn_save.isEnabled()


@pytest.fixture(scope="module")
def qapp():
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    yield app
