"""Monthly birthdays report: month filter, day sorting, xlsx, dialog."""
import os
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from db.export import (
    members_with_birthday_in_month, write_birthday_xlsx, BIRTHDAY_COLUMNS,
)


def member(cid, last, first, dob):
    return {"center_id": cid, "last_name": last, "first_name": first,
            "health_plan": "HF", "dob": dob}


MEMBERS = [
    member(1, "Chan", "Mei", date(1950, 6, 20)),
    member(2, "Lu", "Wei", "1948-06-04"),            # ISO text DOB
    member(3, "Wong", "Ka", "6/4/1952"),             # M/D/YYYY text DOB
    member(4, "Li", "Jun", datetime(1955, 6, 9)),    # terminated below
    member(5, "Ng", "Sam", date(1960, 7, 1)),        # July, not June
    member(6, "Ho", "Yan", None),                    # no DOB -> skipped
    member(7, "Yu", "Lan", "unknown"),               # bad DOB -> skipped
]


def test_month_filter_sorted_by_day_then_name():
    rows = members_with_birthday_in_month(MEMBERS, {4}, 6)
    assert [(r["center_id"], r["dob"]) for r in rows] == [
        (2, date(1948, 6, 4)),          # day 4, "Lu" before "Wong"
        (3, date(1952, 6, 4)),
        (1, date(1950, 6, 20)),
    ]
    assert rows[0]["name"] == "Lu, Wei"


def test_terminated_and_unparseable_excluded():
    rows = members_with_birthday_in_month(MEMBERS, {4}, 6)
    assert 4 not in [r["center_id"] for r in rows]
    assert members_with_birthday_in_month(MEMBERS, set(), 12) == []


def test_xlsx_round_trip(tmp_path):
    from openpyxl import load_workbook
    rows = members_with_birthday_in_month(MEMBERS, set(), 6)
    path = tmp_path / "birthdays.xlsx"
    write_birthday_xlsx(str(path), rows, "June")

    ws = load_workbook(str(path)).active
    assert [c.value for c in ws[1]] == BIRTHDAY_COLUMNS
    assert ws.title == "Birthdays June"
    first = [c.value for c in ws[2]]
    assert first[0] == 2 and first[1] == "Lu, Wei"
    assert first[2].date() == date(1948, 6, 4)
    assert first[3] is None                          # empty Sign column
    assert ws.cell(row=2, column=3).number_format == "MM/DD/YYYY"
    # Same ruled record-sheet treatment as the expiring report.
    for col in range(1, 5):
        assert ws.cell(row=2, column=col).border.bottom.style == "thin"
    assert ws.page_setup.scale == 100
    assert ws.print_title_rows in ("1:1", "$1:$1")


def test_dialog_defaults_and_counts(qapp):
    from gui.birthday_report import BirthdayReportDialog
    dlg = BirthdayReportDialog(MEMBERS, {4}, today=date(2026, 7, 14))
    assert dlg._month.currentText() == "July"        # defaults to this month
    assert "1 member " in dlg._count.text()

    dlg._month.setCurrentIndex(5)                    # June
    assert "3 members" in dlg._count.text()
    assert dlg._btn_save.isEnabled()

    dlg._month.setCurrentIndex(11)                   # December: nobody
    assert "0 members" in dlg._count.text()
    assert not dlg._btn_save.isEnabled()


@pytest.fixture(scope="module")
def qapp():
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    yield app
