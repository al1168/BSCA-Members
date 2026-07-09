"""Member spreadsheet export: row assembly rules and the xlsx round-trip."""
from datetime import date, datetime

from db.export import (
    COLUMNS, build_export_rows, pick_active_auth, format_auth_days_dotted,
    write_members_xlsx,
)

TODAY = date(2026, 7, 9)


def contact(cid, last="Chan", first="Mei", **kw):
    """A Contacts row in _CONTACTS_QUERY column order."""
    fields = {
        "chinese": "陳美", "dob": datetime(1950, 3, 15), "plan": "HF",
        "member_id": "M123", "medicaid": "ab12345c", "medicare": "1eg4te5mk72",
        "ssn": "123456789", "language": "Cantonese", "case_manager": "Lee",
        "home": "2125550100", "cell": "9175550100", "address": "1 Main St",
        "pcp": "Dr. Wu", "hospital": "NYP", "notes": "note", "gender": "F",
        "long_lat": "-73.9,40.7", "hha": "Sunrise",
    }
    fields.update(kw)
    return (cid, last, first, fields["chinese"], fields["dob"], fields["plan"],
            fields["member_id"], fields["medicaid"], fields["medicare"],
            fields["ssn"], fields["language"], fields["case_manager"],
            fields["home"], fields["cell"], fields["address"], fields["pcp"],
            fields["hospital"], fields["notes"], fields["gender"],
            fields["long_lat"], fields["hha"])


def test_column_order_matches_request():
    assert COLUMNS[0] == "Center Id"
    assert COLUMNS[21] == "Enrollment Date"
    assert COLUMNS[22:25] == ["Auth Days", "Auth Start", "Auth End"]
    assert COLUMNS[25:] == ["Emergency_Full Name", "Emergency_Phone Number",
                            "Emergency_Relationship"]


def test_row_is_fully_populated_and_formatted():
    rows = build_export_rows(
        [contact(1)],
        [(1, datetime(2025, 5, 1))],
        [(1, date(2026, 1, 1), date(2026, 12, 31), "1,3,5")],
        [(10, 1, "Wei Lu", "7185550100", "Son")],
        TODAY,
    )
    assert len(rows) == 1
    r = dict(zip(COLUMNS, rows[0]))
    assert r["Center Id"] == 1
    assert r["Chinese Name"] == "陳美"
    assert r["DOB"] == date(1950, 3, 15)          # datetime -> date
    assert r["SSN"] == "123-45-6789"              # display formatting applied
    assert r["Medicaid"] == "AB12345C"
    assert r["Home Tell"] == "(212)-555-0100"
    assert r["Long Lat"] == "-73.9,40.7"
    assert r["Enrollment Date"] == date(2025, 5, 1)
    assert r["Auth Days"] == "1.3.5"
    assert r["Auth Start"] == date(2026, 1, 1)
    assert r["Auth End"] == date(2026, 12, 31)
    assert r["Emergency_Full Name"] == "Wei Lu"
    assert r["Emergency_Phone Number"] == "(718)-555-0100"
    assert r["Emergency_Relationship"] == "Son"


def test_latest_enrollment_start_wins():
    rows = build_export_rows(
        [contact(1)],
        [(1, date(2024, 1, 1)), (1, date(2026, 6, 2)), (1, date(2025, 5, 1))],
        [], [], TODAY)
    assert dict(zip(COLUMNS, rows[0]))["Enrollment Date"] == date(2026, 6, 2)


def test_no_active_auth_leaves_cells_blank():
    rows = build_export_rows(
        [contact(1)],
        [],
        [(1, date(2026, 1, 1), date(2026, 6, 30), "1,2")],   # expired
        [], TODAY)
    r = dict(zip(COLUMNS, rows[0]))
    assert r["Auth Days"] == "" and r["Auth Start"] is None and r["Auth End"] is None


def test_active_auth_selection():
    # Expired, upcoming, and two active — the active one with latest start wins.
    auths = [
        (date(2025, 1, 1), date(2025, 12, 31), "1"),      # expired
        (date(2026, 8, 1), date(2027, 8, 1), "2"),        # upcoming
        (date(2026, 1, 1), date(2026, 12, 31), "3"),      # active
        (date(2026, 6, 1), date(2026, 12, 31), "4"),      # active, later start
    ]
    assert pick_active_auth(auths, TODAY)[2] == "4"
    # Open-ended dates don't disqualify.
    assert pick_active_auth([(None, None, "5")], TODAY)[2] == "5"
    assert pick_active_auth([], TODAY) is None


def test_first_emergency_contact_by_id():
    rows = build_export_rows(
        [contact(1)], [], [],
        [(22, 1, "Second Person", "1", "Friend"),
         (7, 1, "First Person", "2125550100", "Daughter")],
        TODAY)
    r = dict(zip(COLUMNS, rows[0]))
    assert r["Emergency_Full Name"] == "First Person"


def test_all_members_included_even_without_related_rows():
    rows = build_export_rows(
        [contact(1), contact(2, last="Lu", first="Wei")], [], [], [], TODAY)
    assert [r[0] for r in rows] == [1, 2]
    r2 = dict(zip(COLUMNS, rows[1]))
    assert r2["Enrollment Date"] is None and r2["Emergency_Full Name"] == ""


def test_auth_days_dotted():
    assert format_auth_days_dotted("1,3,5") == "1.3.5"
    assert format_auth_days_dotted("3,1,2,4") == "1.2.3.4"   # sorted
    assert format_auth_days_dotted("") == ""
    assert format_auth_days_dotted(None) == ""


def test_xlsx_round_trip(tmp_path):
    from openpyxl import load_workbook
    rows = build_export_rows(
        [contact(1)],
        [(1, date(2025, 5, 1))],
        [(1, date(2026, 1, 1), date(2026, 12, 31), "1,3,5")],
        [(10, 1, "Wei Lu", "7185550100", "Son")],
        TODAY)
    path = tmp_path / "members.xlsx"
    write_members_xlsx(str(path), rows)

    ws = load_workbook(str(path)).active
    header = [c.value for c in ws[1]]
    assert header == COLUMNS
    assert ws.freeze_panes == "A2"
    data = [c.value for c in ws[2]]
    r = dict(zip(header, data))
    assert r["Center Id"] == 1
    assert r["Chinese Name"] == "陳美"
    assert r["Auth Days"] == "1.3.5"
    # Dates come back as datetimes from openpyxl with the display format set.
    assert r["DOB"].date() == date(1950, 3, 15)
    assert ws.cell(row=2, column=5).number_format == "MM/DD/YYYY"
