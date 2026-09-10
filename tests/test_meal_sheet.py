"""Monthly meal sheet: member filter, day shading, xlsx layout, dialog."""
import os
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl.utils import get_column_letter

from db.export import (
    members_for_meal_sheet, closed_days_in_month, meal_day_states,
    write_meal_sheet_xlsx, MEAL_SHEET_TITLES, MEAL_MASK_SHEET,
    MEAL_FIXED_COLUMNS, MEAL_FILL_AUTHORIZED, MEAL_FILL_CLOSED,
    MEAL_FILL_OPEN, DAY_AUTHORIZED, DAY_CLOSED, DAY_NOT_AUTHORIZED,
)


def member(cid, last, first, plan="HF", alt_id=None):
    return {"center_id": cid, "last_name": last, "first_name": first,
            "health_plan": plan, "alt_id": alt_id}


# Deliberately not in center-id order: the sheet must sort by ID itself.
MEMBERS = [
    member(2, "Lu", "Wei", "VCM", 1002),
    member(1, "Chan", "Mei", "HF", 1001),
    member(6, "Ho", "Yan", "AE"),
    member(3, "Wong", "Ka"),          # auth ended before September
    member(4, "Li", "Jun"),           # no enrollment row at all
    member(5, "Ng", "Sam"),           # enrollment ended before September
    member(7, "Tam", "Sze"),          # enrolled, no auth
    member(8, "Yu", "Lin", "SWH"),
    member(9, "Fan", "Bo"),           # enrollment with no start date
]

GROUPS = {1: "C", 2: "B", 8: "C", 3: "B"}      # 6 missing -> blank Location

# (center_id, start_date, end_date) — as _ENROLLMENTS_QUERY returns them.
ENROLLMENTS = [
    (1, date(2025, 3, 1), None),
    (2, date(2025, 1, 1), date(2026, 9, 5)),       # ends inside September
    (3, date(2025, 1, 1), None),
    (5, date(2025, 1, 1), date(2026, 8, 15)),      # ended before September
    (6, datetime(2026, 9, 28), None),              # starts inside September
    (7, date(2025, 1, 1), None),
    (8, date(2026, 1, 1), None),
    (9, None, None),                               # no start date
    (None, date(2025, 1, 1), None),                # no Center ID
]

# (center_id, auth_start, auth_end, auth_days) — as _AUTHS_QUERY returns them.
AUTHS = [
    (1, date(2026, 9, 10), date(2026, 9, 20), "1,3"),
    (2, date(2026, 1, 1), date(2026, 12, 31), "1"),
    (2, date(2026, 6, 1), date(2026, 12, 31), "5"),
    (3, date(2025, 9, 1), date(2026, 8, 31), "1,2,3"),
    (4, date(2026, 1, 1), date(2026, 12, 31), "1"),
    (5, date(2026, 1, 1), date(2026, 12, 31), "1"),
    (6, datetime(2026, 9, 1), None, "1,2,3,4,5,6,7"),
    (8, None, date(2026, 12, 31), "2"),
    (9, date(2026, 1, 1), date(2026, 12, 31), "1"),
    (None, date(2026, 1, 1), date(2026, 12, 31), "1"),
]

HOLIDAYS = [
    {"id": 1, "name": "Labor Day", "date": date(2026, 9, 7)},
    {"id": 2, "name": "Thanksgiving", "date": date(2026, 11, 26)},
]
OPERATING = {d: {"id": d, "day_of_week": d} for d in range(1, 6)}   # Mon–Fri
CLOSED_SEPT = {5, 6, 7, 12, 13, 19, 20, 26, 27}

INPUTS = {"enrollments": ENROLLMENTS, "auths": AUTHS, "groups": GROUPS,
          "holidays": HOLIDAYS, "operating_days": OPERATING}


def rows_for(year=2026, month=9, unlocked=True):
    return members_for_meal_sheet(MEMBERS, GROUPS, ENROLLMENTS, AUTHS,
                                  year, month, unlocked)


# ── member filter ────────────────────────────────────────────────────────

def test_requires_enrollment_and_auth_overlap():
    ids = [r["center_id"] for r in rows_for()]
    assert ids == [1, 2, 6, 8]
    # 3: auth ended Aug 31; 4: never enrolled; 5: enrollment ended Aug 15;
    # 7: no auth; 9: enrollment has no start date.


def test_row_fields_and_sort():
    rows = rows_for()
    assert rows[0]["name"] == "Chan, Mei"
    assert rows[0]["health_plan"] == "HF"
    assert rows[0]["location"] == "C"
    assert rows[0]["alt_id"] == 1001
    assert rows[2]["location"] == ""              # 6 has no Group
    assert rows[2]["center_id"] == 6
    assert rows[1]["auths"] == [
        (date(2026, 1, 1), date(2026, 12, 31), "1"),
        (date(2026, 6, 1), date(2026, 12, 31), "5"),
    ]
    # Access datetimes become plain dates; None start/end pass through.
    assert rows[2]["auths"] == [(date(2026, 9, 1), None, "1,2,3,4,5,6,7")]
    assert rows[3]["auths"] == [(None, date(2026, 12, 31), "2")]


def test_alt_id_blank_when_locked():
    assert all(r["alt_id"] is None for r in rows_for(unlocked=False))


def test_other_months():
    # May: 3's auth (to Aug 31) and 5's enrollment (to Aug 15) still overlap.
    assert [r["center_id"] for r in rows_for(2026, 5)] == [2, 3, 5, 8]
    # October: 1's auth and 2's enrollment ended in September.
    assert [r["center_id"] for r in rows_for(2026, 10)] == [6, 8]
    assert rows_for(2024, 5) == []


# ── closed days ──────────────────────────────────────────────────────────

def test_closed_days_in_month():
    assert closed_days_in_month(2026, 9, HOLIDAYS, OPERATING) == CLOSED_SEPT
    # Weekend-only when there are no holidays that month.
    assert closed_days_in_month(2026, 10, HOLIDAYS, OPERATING) == {
        3, 4, 10, 11, 17, 18, 24, 25, 31}
    # No operating rows at all: every day is closed.
    assert closed_days_in_month(2026, 9, [], {}) == set(range(1, 31))


# ── day states ───────────────────────────────────────────────────────────

def test_meal_day_states_range_and_weekday():
    states = meal_day_states([(date(2026, 9, 10), date(2026, 9, 20), "1,3")],
                             2026, 9, CLOSED_SEPT)
    assert len(states) == 30
    assert states[13] == DAY_AUTHORIZED          # Mon 14
    assert states[15] == DAY_AUTHORIZED          # Wed 16
    assert states[8] == DAY_NOT_AUTHORIZED       # Wed 9, before start
    assert states[20] == DAY_NOT_AUTHORIZED      # Mon 21, after end
    assert states[5] == DAY_CLOSED               # Sun 6
    assert states[6] == DAY_CLOSED               # Labor Day


def test_meal_day_states_closed_wins_and_union():
    both = [(date(2026, 9, 1), date(2026, 9, 30), "1"),
            (date(2026, 9, 1), date(2026, 9, 30), "5")]
    states = meal_day_states(both, 2026, 9, CLOSED_SEPT)
    assert states[6] == DAY_CLOSED               # Mon 7 is a holiday
    assert states[13] == DAY_AUTHORIZED          # Mon 14
    assert states[17] == DAY_AUTHORIZED          # Fri 18
    assert states[14] == DAY_NOT_AUTHORIZED      # Tue 15


def test_meal_day_states_open_ended_and_malformed():
    states = meal_day_states([(None, None, "2")], 2026, 9, set())
    assert [i + 1 for i, s in enumerate(states) if s == DAY_AUTHORIZED] == [
        1, 8, 15, 22, 29]
    states = meal_day_states([(None, None, "Mon")], 2026, 9, set())
    assert DAY_AUTHORIZED not in states
    assert meal_day_states([], 2026, 9, set()) == [DAY_NOT_AUTHORIZED] * 30


# ── workbook ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    from openpyxl import load_workbook
    path = tmp_path_factory.mktemp("meal") / "meal.xlsx"
    write_meal_sheet_xlsx(str(path), rows_for(), 2026, 9, CLOSED_SEPT)
    return load_workbook(str(path))


LAST_DAY, TOTAL, LOCATION = 34, 35, 36       # 30-day month: AH, AI, AJ


def test_xlsx_sheets_and_header(workbook):
    wb = workbook
    assert wb.sheetnames == list(MEAL_SHEET_TITLES) + [MEAL_MASK_SHEET]
    assert wb[MEAL_MASK_SHEET].sheet_state == "hidden"
    assert wb.active.title == MEAL_SHEET_TITLES[0]
    ws = wb[MEAL_SHEET_TITLES[0]]
    assert [c.value for c in ws[1]][:4] == MEAL_FIXED_COLUMNS
    assert ws.cell(1, 5).value.date() == date(2026, 9, 1)
    assert ws.cell(1, LAST_DAY).value.date() == date(2026, 9, 30)
    assert ws.cell(1, 5).number_format == "m/d;@"
    assert ws.cell(1, TOTAL).value == "Total"
    assert ws.cell(1, LOCATION).value == "Location"
    assert ws.cell(1, LOCATION + 1).value is None
    assert ws["A1"].font.bold and ws["A1"].font.size == 14
    assert ws["A1"].font.name == "Calibri"
    assert ws.row_dimensions[1].height == 44
    assert ws.row_dimensions[2].height == 26
    widths = {k: ws.column_dimensions[k].width for k in ("A", "B", "C", "D")}
    assert widths == {"A": 16, "B": 16, "C": 29, "D": 18}
    assert ws.column_dimensions["E"].width == 9
    assert ws.column_dimensions[get_column_letter(TOTAL)].width == 9
    assert ws.column_dimensions[get_column_letter(LOCATION)].width == 10
    assert ws.freeze_panes == "E2"
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.scale == 79
    assert ws.print_title_rows in ("1:1", "$1:$1")
    assert ws.print_title_cols in ("A:D", "$A:$D")
    assert ws.page_margins.left == 0.25 and ws.page_margins.top == 0.75
    for ref in ("A1", "E1", "A2", "E2", "AI2", "AJ2", "A6", "E8"):
        assert ws[ref].border.bottom.style == "thin", ref


def test_xlsx_member_rows(workbook):
    ws = workbook[MEAL_SHEET_TITLES[0]]
    assert [ws.cell(r, 1).value for r in range(2, 6)] == [1, 2, 6, 8]
    assert ws["B2"].value == 1001
    assert ws["B4"].value is None                 # member 6 has no alt id
    assert ws["C2"].value == "Chan, Mei"
    assert ws["D2"].value == "HF"
    assert ws.cell(2, LOCATION).value == "C"
    assert ws.cell(4, LOCATION).value in (None, "")
    assert ws.cell(2, 5).value is None            # day cells are blank
    assert ws.cell(2, TOTAL).value == "=SUM(E2:AH2)"


def test_xlsx_fills_and_mask(workbook):
    ws = workbook[MEAL_SHEET_TITLES[0]]
    mask = workbook[MEAL_MASK_SHEET]
    sat5, wed9, mon14 = 9, 13, 18                # columns of Sep 5 / 9 / 14
    assert ws.cell(2, sat5).fill.fgColor.rgb.endswith(MEAL_FILL_CLOSED)
    assert ws.cell(2, mon14).fill.fgColor.rgb.endswith(MEAL_FILL_AUTHORIZED)
    assert ws.cell(2, wed9).fill.fgColor.rgb.endswith(MEAL_FILL_OPEN)
    assert mask.cell(2, mon14).value == 1
    assert mask.cell(2, wed9).value == 0
    assert mask.cell(2, sat5).value == 0
    # Member 6 (row 4) is authorized every day, so only closed days differ.
    assert ws.cell(4, 4 + 28).fill.fgColor.rgb.endswith(MEAL_FILL_AUTHORIZED)
    assert ws.cell(4, 4 + 27).fill.fgColor.rgb.endswith(MEAL_FILL_CLOSED)
    # Member 8 (row 5): Tuesdays only.
    assert ws.cell(5, 4 + 15).fill.fgColor.rgb.endswith(MEAL_FILL_AUTHORIZED)
    assert ws.cell(5, 4 + 16).fill.fgColor.rgb.endswith(MEAL_FILL_OPEN)
    assert mask.cell(5, 4 + 15).value == 1 and mask.cell(5, 4 + 16).value == 0


def test_xlsx_summary_rows(workbook):
    ws = workbook[MEAL_SHEET_TITLES[0]]
    assert ws["A6"].value == "Total"
    assert ws["E6"].value == "=SUM(E2:E5)"
    assert ws.cell(6, LAST_DAY).value == "=SUM(AH2:AH5)"
    assert ws["A7"].value == "Authorized"
    assert ws["E7"].value == "=SUMPRODUCT(E2:E5,'_mask'!E2:E5)"
    assert ws["A8"].value == "Not_Authorized"
    assert ws["E8"].value == "=SUMPRODUCT(E2:E5,1-'_mask'!E2:E5)"
    for r in (6, 7, 8):
        assert ws.cell(r, TOTAL).value == f"=SUM(E{r}:AH{r})"
    assert ws["A9"].value is None


def test_xlsx_three_sheets_identical(workbook):
    first = workbook[MEAL_SHEET_TITLES[0]]
    for title in MEAL_SHEET_TITLES[1:]:
        ws = workbook[title]
        for r in (1, 3):
            assert [ws.cell(r, c).value for c in range(1, LOCATION + 1)] ==                 [first.cell(r, c).value for c in range(1, LOCATION + 1)]
        assert ws.cell(2, 18).fill.fgColor.rgb == first.cell(2, 18).fill.fgColor.rgb
        assert ws.freeze_panes == "E2"
        assert ws["E7"].value == first["E7"].value


def test_xlsx_zero_rows_header_only(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "empty.xlsx"
    write_meal_sheet_xlsx(str(path), [], 2026, 2, set())
    ws = load_workbook(str(path))[MEAL_SHEET_TITLES[0]]
    assert ws.max_row == 1
    assert ws.cell(1, 4 + 28).value.date() == date(2026, 2, 28)
    assert ws.cell(1, 4 + 29).value == "Total"


# ── dialog ───────────────────────────────────────────────────────────────

def test_dialog_defaults_and_counts(qapp, monkeypatch):
    from gui import meal_sheet as ms
    calls = []
    monkeypatch.setattr("db.export.load_meal_sheet_inputs",
                        lambda db: calls.append(db) or INPUTS)
    dlg = ms.MealSheetDialog("fake.accdb", MEMBERS, today=date(2026, 9, 15))
    assert dlg._month.currentText() == "September"
    assert dlg._year.value() == 2026
    assert "4 members" in dlg._count.text()
    assert dlg._btn_save.isEnabled()

    dlg._month.setCurrentIndex(9)                    # October 2026
    assert "2 members" in dlg._count.text()

    dlg._year.setValue(2024)
    assert "0 members" in dlg._count.text()
    assert not dlg._btn_save.isEnabled()
    assert calls == ["fake.accdb"]                   # inputs loaded once


def test_dialog_load_error_disables_save(qapp, monkeypatch):
    from gui import meal_sheet as ms
    shown = []
    monkeypatch.setattr("gui.errors.show_db_error",
                        lambda parent, exc: shown.append(str(exc)))

    def boom(db):
        raise RuntimeError("cannot find the input table 'Holidays'")
    monkeypatch.setattr("db.export.load_meal_sheet_inputs", boom)
    dlg = ms.MealSheetDialog("fake.accdb", MEMBERS, today=date(2026, 9, 15))
    assert shown and "Holidays" in shown[0]
    assert not dlg._btn_save.isEnabled()


def test_dialog_save_writes_workbook(qapp, monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from gui import meal_sheet as ms
    monkeypatch.setattr("db.export.load_meal_sheet_inputs", lambda db: INPUTS)
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(ms.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(out), "")))
    monkeypatch.setattr(ms.QMessageBox, "exec", lambda self: 0)
    dlg = ms.MealSheetDialog("fake.accdb", MEMBERS, today=date(2026, 9, 15),
                             alt_ids_unlocked=True)
    dlg._save()
    ws = load_workbook(str(out))[MEAL_SHEET_TITLES[2]]
    assert ws["A2"].value == 1 and ws["B2"].value == 1001
    assert ws.cell(2, 11).fill.fgColor.rgb.endswith(MEAL_FILL_CLOSED)  # Sep 7


# ── main-window wiring ───────────────────────────────────────────────────

def test_toolbar_button_sits_between_calendar_and_export(qapp, tmp_path):
    from PyQt6.QtWidgets import QToolBar
    from gui.main_window import MainWindow
    w = MainWindow({"db_path": "", "theme": "dark"},
                   str(tmp_path / "settings.json"))
    toolbar = w.findChild(QToolBar, "main_toolbar")
    names = [toolbar.widgetForAction(a).objectName()
             for a in toolbar.actions()
             if toolbar.widgetForAction(a) is not None]
    i = names.index("btn_meal_sheet")
    assert names[i - 1] == "btn_company_calendar"
    assert names[i + 1] == "btn_export"


def test_open_meal_sheet_without_db_warns(qapp, tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QMessageBox
    from gui.main_window import MainWindow
    warned = []
    monkeypatch.setattr(QMessageBox, "warning",
                        lambda *a, **k: warned.append(a[1]))
    w = MainWindow({"db_path": "", "theme": "dark"},
                   str(tmp_path / "settings.json"))
    w._open_meal_sheet()
    assert warned == ["No Database"]


@pytest.fixture(scope="module")
def qapp():
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    yield app
