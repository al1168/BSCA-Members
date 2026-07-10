"""Monthly expiring-auths report: month filter, plan grouping, xlsx, HTML."""
import os
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from db.export import members_expiring_in_month, write_expiring_xlsx, EXPIRING_COLUMNS
from gui.expiring_report import build_expiring_report_html


def member(cid, last, first, plan):
    return {"center_id": cid, "last_name": last, "first_name": first,
            "health_plan": plan}


MEMBERS = [
    member(1, "Chan", "Mei", "HF"),
    member(2, "Lu", "Wei", "AE"),
    member(3, "Wong", "Ka", "HF"),
    member(4, "Li", "Jun", "AE"),      # terminated below
    member(5, "Ng", "Sam", "VCM"),     # no auths
]


def test_month_filter_and_plan_grouping():
    auths = [
        (1, date(2026, 8, 20)),
        (2, date(2026, 8, 5)),
        (3, date(2026, 8, 1)),
        (4, date(2026, 8, 9)),         # terminated -> excluded
        (2, date(2026, 1, 1)),         # older auth ignored (latest end wins)
    ]
    rows = members_expiring_in_month(MEMBERS, {4}, auths, 2026, 8)
    # Grouped by plan (AE then HF), soonest end first within a plan.
    assert [(r["health_plan"], r["center_id"]) for r in rows] == \
        [("AE", 2), ("HF", 3), ("HF", 1)]
    assert rows[0]["name"] == "Lu, Wei"
    assert rows[0]["end"] == date(2026, 8, 5)


def test_renewal_pushes_member_out_of_month():
    auths = [(1, date(2026, 8, 20)), (1, date(2027, 8, 20))]  # renewed
    assert members_expiring_in_month(MEMBERS, set(), auths, 2026, 8) == []


def test_wrong_year_not_included():
    auths = [(1, date(2027, 8, 20))]
    assert members_expiring_in_month(MEMBERS, set(), auths, 2026, 8) == []


def test_datetime_ends_normalized():
    auths = [(1, datetime(2026, 8, 20, 0, 0))]
    rows = members_expiring_in_month(MEMBERS, set(), auths, 2026, 8)
    assert rows[0]["end"] == date(2026, 8, 20)


def test_xlsx_round_trip(tmp_path):
    from openpyxl import load_workbook
    rows = members_expiring_in_month(
        MEMBERS, set(),
        [(1, date(2026, 8, 20)), (2, date(2026, 8, 5))], 2026, 8)
    path = tmp_path / "expiring.xlsx"
    write_expiring_xlsx(str(path), rows, "August 2026")

    ws = load_workbook(str(path)).active
    assert [c.value for c in ws[1]] == EXPIRING_COLUMNS
    assert ws.freeze_panes == "A2"
    first = [c.value for c in ws[2]]
    assert first[0] == 2 and first[1] == "Lu, Wei" and first[2] == "AE"
    assert first[3].date() == date(2026, 8, 5)
    assert first[4] is None                    # empty Notes column
    assert ws.cell(row=2, column=4).number_format == "MM/DD/YYYY"

    # A record sheet: every cell ruled, including the empty Notes boxes;
    # header shaded and repeated on each printed page; fit-to-width printing.
    for col in range(1, 6):
        assert ws.cell(row=1, column=col).border.top.style == "thin"
        assert ws.cell(row=2, column=col).border.bottom.style == "thin"
    assert ws.cell(row=2, column=5).border.left.style == "thin"
    assert ws.cell(row=1, column=1).fill.fgColor.rgb.endswith("E8E8E8")
    assert ws.print_title_rows in ("1:1", "$1:$1")
    assert ws.page_setup.fitToWidth == 1
    assert ws.row_dimensions[2].height == 24


def test_html_has_rows_notes_column_and_title():
    rows = members_expiring_in_month(
        MEMBERS, set(),
        [(1, date(2026, 8, 20)), (2, date(2026, 8, 5))], 2026, 8)
    html = build_expiring_report_html(rows, "August 2026", "07/10/2026")
    assert "Expiring Authorizations — August 2026" in html
    assert "Lu, Wei" in html and "Chan, Mei" in html
    assert "08/05/2026" in html
    assert "Notes" in html
    assert "2 members" in html


def test_html_empty_month_message():
    html = build_expiring_report_html([], "March 2026", "07/10/2026")
    assert "No members have authorizations expiring" in html


def test_dialog_counts_and_defaults(qapp, monkeypatch):
    from gui import expiring_report as er
    monkeypatch.setattr(
        "db.members.get_member_auth_ends",
        lambda _db: [(1, date(2026, 8, 20)), (2, date(2026, 8, 5))])
    dlg = er.ExpiringReportDialog(
        "fake.accdb", MEMBERS, set(), today=date(2026, 7, 10))
    assert dlg._month.currentText() == "August"       # defaults to next month
    assert dlg._year.value() == 2026
    assert "2 members" in dlg._count.text()
    assert dlg._btn_save.isEnabled() and dlg._btn_print.isEnabled()

    dlg._month.setCurrentIndex(2)                     # March: nothing expiring
    assert "0 members" in dlg._count.text()
    assert not dlg._btn_save.isEnabled()


@pytest.fixture(scope="module")
def qapp():
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    yield app
